"""Independent source readers and calculations; never imports production code.

This is a QA reference for the received synthetic case, not accounting advice.
Usage: python -X utf8 tests/independent/reference.py --output <json>
"""
import argparse
import csv
import hashlib
import io
import json
import re
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP
from email import policy
from email.parser import BytesParser
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "contexto/data_pack"
CENT = Decimal("0.01")


def br(value):
    return Decimal(str(value).replace(".", "").replace(",", "."))


def rounded(value):
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def cnpj(value):
    return re.sub(r"\D", "", str(value)).zfill(14)


def rate(value):
    if isinstance(value, str):
        return br(value.replace("%", "")) / 100
    return Decimal(str(value))


def table(path, encoding="utf-8-sig"):
    with path.open(encoding=encoding, newline="") as stream:
        return list(csv.DictReader(stream, delimiter=";"))


def read_sources():
    checkout = table(DATA / "01_orbe_relatorio_comissoes_set2026.csv")
    workbook = openpyxl.load_workbook(DATA / "02_tabela_taxas_contratadas.xlsx", data_only=True)
    contracts = [(n, list(row)) for n, row in enumerate(workbook.worksheets[0].iter_rows(values_only=True), 1)
                 if 6 <= n <= 30]
    message = BytesParser(policy=policy.default).parsebytes((DATA / "03_email_ponte_cambio_relatorio_set2026.eml").read_bytes())
    attachment = next(part for part in message.iter_attachments() if part.get_filename().lower().endswith(".xlsx"))
    ponte_book = openpyxl.load_workbook(io.BytesIO(attachment.get_payload(decode=True)), data_only=True)
    remittances = [(n, list(row)) for n, row in enumerate(ponte_book["Operações"].iter_rows(values_only=True), 1)
                   if isinstance(row[0], str) and row[0].startswith("PC-")]
    banks = []
    aurora = DATA / "05a_extrato_banco_aurora_cc4471_set2026.csv"
    for n, row in enumerate(csv.reader(aurora.read_text(encoding="utf-8-sig").splitlines(), delimiter=";"), 1):
        if len(row) >= 4 and re.fullmatch(r"\d{2}/09/2026", row[0]) and row[3]:
            banks.append({"bank": "AURORA", "ref": f"L{n}", "description": row[1], "value": str(br(row[3])), "file": aurora.name})
    meridiano = DATA / "05b_extrato_banco_meridiano_cc90233_set2026.ofx"
    for block in re.findall(r"<STMTTRN>(.*?)</STMTTRN>", meridiano.read_text(encoding="cp1252"), re.S):
        get = lambda tag: re.search(fr"<{tag}>([^<\r\n]+)", block)[1].strip()
        banks.append({"bank": "MERIDIANO", "ref": get("FITID"), "description": get("MEMO"), "value": get("TRNAMT"), "file": meridiano.name})
    return checkout, contracts, remittances, banks


def build_reference():
    checkout, contracts, remittances, banks = read_sources()
    by_ec = defaultdict(list)
    mappings = {"1043300981": "66564512000199", "1043396662": "78576234000179", "1038867654": "07752116000126"}
    for physical, contract in contracts:
        if contract[2]:
            by_ec[str(contract[2])].append((physical, contract))
        for ec, identity in mappings.items():
            if cnpj(contract[1]) == identity:
                by_ec[ec].append((physical, contract))
    expected_checkout = {}
    atlas_refs, historical_refs = [], []
    by_id = {r["id_transacao"]: r for r in checkout}
    for row in checkout:
        ref = row["id_transacao"]
        ec = row["cod_estabelecimento"]
        if ec == "1076801732":
            atlas_refs.append(ref)
            continue
        effective = row["data_venda_original"] or row["data_movimento"]
        year_month = effective[6:] + "-" + effective[3:5]
        historical = year_month < "2026-09"
        if historical:
            historical_refs.append(ref)
        candidates = by_ec[ec]
        contract_line, contract = candidates[0]
        if len(candidates) > 1 and (year_month > "2026-09" or year_month == "2026-09" and int(effective[:2]) >= 15):
            contract_line, contract = candidates[-1]
        kind = row["tipo_lancamento"]
        if "ANTECIPACAO" in kind:
            col = 8
        else:
            col = {"DEBITO": 3, "CREDITO A VISTA": 4, "PIX": 7}.get(row["modalidade"])
            if col is None:
                col = 5 if int(row["qtd_parcelas"]) <= 6 else 6
        contracted = rate(contract[col])
        if ec == "1018325654" and "2026-09" <= year_month <= "2026-12":
            contracted = Decimal(0)
        value = rounded(br(row["valor_bruto"]) * contracted)
        applied = br(row["taxa_aplicada_pct"])
        expected_checkout[ref] = {"cnpj": cnpj(contract[1]), "value": str(value), "contract_rate_pct": str(contracted * 100),
                                  "applied_rate_pct": str(applied), "difference": str(rounded(br(row["valor_bruto"]) * (applied / 100 - contracted))),
                                  "contract_row": contract_line, "effective": effective,
                                  "confidence": "historical_proxy_requires_review" if historical else "source_supported",
                                  "original_present": row["id_referencia"] in by_id if row["id_referencia"] else None}
    remit_totals = defaultdict(Decimal)
    remit_contracts = {cnpj(row[1]): (Decimal("0.015") if row[9] in {None, "padrão"} else rate(row[9])) for _, row in contracts}
    remit_lines, duplicates, negatives = {}, [], []
    for physical, row in remittances:
        ref = row[0]
        if ref in remit_lines:
            duplicates.append({"ref": ref, "first_row": remit_lines[ref]["physical_row"], "duplicate_row": physical})
            continue
        spread = Decimal(str(row[10]))
        eligible = row[1].strftime("%Y-%m") == "2026-09" and spread >= 0
        share = rounded(spread / 2) if eligible else Decimal(0)
        remit_lines[ref] = {"physical_row": physical, "cnpj": cnpj(row[3]), "spread": str(spread), "value": str(share),
                           "applied_rate_pct": str(Decimal(str(row[11])) * 100), "contract_rate_pct": str(remit_contracts[cnpj(row[3])] * 100),
                           "difference": str(rounded(Decimal(str(row[6])) * Decimal(str(row[8])) * (Decimal(str(row[11])) - remit_contracts[cnpj(row[3])]))),
                           "eligible": eligible, "contract_date": row[1].isoformat(), "settle_date": row[2].isoformat(),
                           "status": row[14]}
        if spread < 0:
            negatives.append(ref)
        if eligible:
            remit_totals[cnpj(row[3])] += share
    files = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in DATA.iterdir() if p.is_file()}
    return {"reference_version": "QA-R004-v1", "period": "2026-09", "source_hashes": files,
            "scope": "received pack only; no late adjustment file", "late_adjustments_present": (DATA / "06_orbe_ajustes_tardios_set2026.csv").exists(),
            "counts": {"checkout_rows": len(checkout), "checkout_unique_ids": len(by_id), "remittance_rows": len(remittances),
                       "remittance_unique_ids": len(remit_lines), "banks": dict(Counter(b["bank"] for b in banks))},
            "checkout_lines": expected_checkout, "atlas_refs": atlas_refs, "historical_proxy_refs": historical_refs,
            "remittance_lines": remit_lines, "remittance_totals": {k: str(v) for k, v in remit_totals.items()},
            "remittance_total": str(sum(remit_totals.values(), Decimal(0))), "duplicates": duplicates, "negative_refs": negatives, "banks": banks,
            "limits": ["Historical checkout values are review proxies, not independent proof of missing historical contracts.",
                       "No official answer key, source correctness or real financial approval is certified.",
                       "No tolerance for remittance divergence was imposed by client; assess implementation rationale separately."]}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = build_reference()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: result[k] for k in ["reference_version", "counts", "late_adjustments_present", "remittance_total", "duplicates", "negative_refs"]}, ensure_ascii=False))
